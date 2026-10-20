"""名单导出：写成 Excel/CSV 给人看或交给别人补数据。

默认导出 .xlsx 而不是 .csv：
  - CSV 用 UTF-8 存，中文 Excel 双击打开是乱码，得靠 BOM 兜着
  - 更严重的是长号码（身份证号）在 CSV 里会被 Excel 读成数字，
    超过 15 位有效数字后末尾几位永久丢失，用户完全看不出来
  - xlsx 可以把号码列显式设为文本格式，从源头避免这个损坏

openpyxl 已经是项目依赖，没有额外成本。
"""

import datetime
import os
from typing import List, Optional, Sequence

from core.roster import Roster, STATUS_GENERATED

# 这些列一律按文本写，避免 Excel 把长号码转成数字或科学计数法
TEXT_COLUMNS_HINT = (
    "身份证", "证件", "卡号", "账号", "账户", "手机", "电话",
    "编号", "号码", "流水", "社保", "公积金",
)

# 表头样式
HEADER_FILL = "FFDCE6F1"
HEADER_FONT_COLOR = "FF1F2937"
PENDING_FILL = "FFFFF7CC"      # 待补充的格子：淡黄，和主界面"黄色=还没填"一致
GENERATED_FILL = "FFE8F5E9"     # 已生成的行：淡绿
AMBIGUOUS_FILL = "FFFDE2E2"     # 重名行：淡红，提醒用户核对


def looks_like_long_number_column(column: str) -> bool:
    """判断列名是否可能装长号码，需要按文本写。"""
    text = str(column or "")
    return any(hint in text for hint in TEXT_COLUMNS_HINT)


def default_file_name(business_name: str, suffix: str = "") -> str:
    """生成默认导出文件名：名单_业务名_年月[_后缀].xlsx"""
    stamp = datetime.date.today().strftime("%Y%m")
    name = f"名单_{business_name or '未命名业务'}_{stamp}"
    if suffix:
        name += f"_{suffix}"
    return f"{name}.xlsx"


def export_to_xlsx(roster: Roster, path: str,
                   columns: Optional[Sequence[str]] = None,
                   include_status: bool = True,
                   required: Optional[Sequence[str]] = None) -> int:
    """把名单导出为 Excel 文件，返回导出的数据行数。

    :param include_status: 是否附加"状态/缺失项"两列，方便对方知道谁还没填完
    :param required: 必填列名，用于算出每行还缺什么；不传则只显示状态
    """
    try:
        import openpyxl
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter
    except ImportError:
        raise RuntimeError("缺少 openpyxl，无法导出 Excel 文件。")

    header = [str(c) for c in (columns or roster.columns)]
    if not header:
        raise RuntimeError("名单没有任何列，无法导出。")

    required_fields = [str(name) for name in (required or [])]

    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "人员名单"

    # 表头
    full_header = list(header)
    if include_status:
        full_header += ["状态", "还缺什么"]
    sheet.append(full_header)

    header_font = Font(bold=True, color=HEADER_FONT_COLOR)
    header_fill = PatternFill("solid", fgColor=HEADER_FILL)
    for index in range(1, len(full_header) + 1):
        cell = sheet.cell(row=1, column=index)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")

    # 数据行：号码列强制文本格式
    text_columns = {index + 1 for index, name in enumerate(header)
                    if looks_like_long_number_column(name)}

    for row_index, row in enumerate(roster.rows, start=2):
        values = [str(row.values.get(name, "") or "") for name in header]
        missing = row.missing_fields(required_fields) if required_fields else []
        if include_status:
            values.append(_status_text(row, bool(missing)))
            values.append("、".join(missing))

        row_fill = _row_fill(row, missing)
        for column_index, value in enumerate(values, start=1):
            cell = sheet.cell(row=row_index, column=column_index, value=value)
            if column_index in text_columns:
                # 关键：显式设成文本，否则 Excel 会把 18 位号码当数字截断
                cell.number_format = "@"
            if row_fill is not None:
                cell.fill = row_fill

    # 列宽按内容自适应，上限 40 避免一列撑满屏幕
    for column_index, name in enumerate(full_header, start=1):
        width = _display_width(str(name)) + 4
        for row in sheet.iter_rows(min_row=2, min_col=column_index,
                                   max_col=column_index, values_only=True):
            width = max(width, _display_width(str(row[0] or "")) + 2)
        sheet.column_dimensions[get_column_letter(column_index)].width = min(width, 40)

    sheet.freeze_panes = "A2"       # 冻结表头，滚动时始终能看到列名
    sheet.auto_filter.ref = f"A1:{get_column_letter(len(full_header))}{sheet.max_row}"

    folder = os.path.dirname(path)
    if folder:
        os.makedirs(folder, exist_ok=True)
    workbook.save(path)
    return len(roster.rows)


def export_to_csv(roster: Roster, path: str,
                  columns: Optional[Sequence[str]] = None) -> int:
    """导出为 CSV（UTF-8 with BOM，Excel 双击不乱码）。

    注意：CSV 无法保存单元格格式，长号码在 Excel 里重新打开时
    仍可能被当成数字。要给别人补数据请优先用 Excel 导出。
    """
    import csv

    header = [str(c) for c in (columns or roster.columns)]
    if not header:
        raise RuntimeError("名单没有任何列，无法导出。")

    folder = os.path.dirname(path)
    if folder:
        os.makedirs(folder, exist_ok=True)

    with open(path, "w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.writer(handle)
        writer.writerow(header)
        for row in roster.rows:
            writer.writerow([str(row.values.get(name, "") or "") for name in header])
    return len(roster.rows)


def write_template_xlsx(path: str, columns: Sequence[str],
                        key_field: str = "") -> None:
    """导出一份空模板（只有表头 + 一行示例），供用户按格式准备数据。

    示例行特意写成文本格式的号码，让用户照抄时不会踩精度坑。
    """
    try:
        import openpyxl
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter
    except ImportError:
        raise RuntimeError("缺少 openpyxl，无法生成模板文件。")

    names = [str(c) for c in columns if str(c).strip()]
    if not names:
        raise RuntimeError("没有可用的列名，无法生成模板。")

    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "名单模板"
    sheet.append(names)

    header_font = Font(bold=True, color=HEADER_FONT_COLOR)
    header_fill = PatternFill("solid", fgColor=HEADER_FILL)
    for index in range(1, len(names) + 1):
        cell = sheet.cell(row=1, column=index)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")

    # 示例行：号码列用文本格式，提示用户该怎么填
    sample = []
    for name in names:
        if looks_like_long_number_column(name):
            sample.append("360700199001011234")
        elif name == key_field or "姓名" in name:
            sample.append("张三")
        else:
            sample.append("")
    sheet.append(sample)
    for index, name in enumerate(names, start=1):
        if looks_like_long_number_column(name):
            sheet.cell(row=2, column=index).number_format = "@"

    for index, name in enumerate(names, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = max(
            _display_width(name) + 6, 14)
    sheet.freeze_panes = "A2"

    note = sheet.cell(row=4, column=1,
                      value="说明：第一行是变量标题，从第二行开始每行一个人。"
                            "身份证/卡号等长号码请把单元格格式设为「文本」再输入，"
                            "否则 Excel 会丢失末尾几位数字。")
    note.font = Font(color="FF6B7280", italic=True)

    folder = os.path.dirname(path)
    if folder:
        os.makedirs(folder, exist_ok=True)
    workbook.save(path)


def _row_fill(row, missing: List[str]):
    """按行状态返回底色，不需要上色时返回 None。

    待补充=淡黄（和主界面"黄色格子=还没填"一致），已生成=淡绿。
    导出文件常要交给别人补数据，上色能让对方一眼看出还差谁。
    """
    from openpyxl.styles import PatternFill

    if row.status == STATUS_GENERATED:
        return PatternFill("solid", fgColor=GENERATED_FILL)
    if missing:
        return PatternFill("solid", fgColor=PENDING_FILL)
    return None


def _status_text(row, has_missing: bool) -> str:
    """状态列的中文文案。"""
    if row.status == STATUS_GENERATED:
        stamp = f"（{row.generated_at}）" if row.generated_at else ""
        return f"已生成{stamp}"
    return "待补充" if has_missing else "已填完"


def _display_width(text: str) -> int:
    """估算显示宽度：中文算 2 个字符宽，用于设置列宽。"""
    width = 0
    for char in str(text or ""):
        width += 2 if ord(char) > 127 else 1
    return width
