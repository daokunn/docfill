"""
全面测试脚本 - 作为专业测试工程师对系统进行压力测试
覆盖：parser / scanner / renderer / variable_manager / output / 全流程
"""

import os
import sys
import traceback

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from core.parser import SmartParser
from core.scanner import scan_businesses, extract_template_variables
from core.renderer import render_document, RenderError
from core.variable_manager import VariableManager
from core.expression import evaluate, try_evaluate, is_expression, ExpressionError
from output.file_output import OutputManager, OutputError

TEST_TEMPLATE_ROOT = os.path.join(os.path.dirname(os.path.abspath(__file__)), "comprehensive_templates")
TEST_OUTPUT_ROOT = os.path.join(os.path.dirname(os.path.abspath(__file__)), "comprehensive_output")

# 测试统计
passed = 0
failed = 0
bugs = []  # (模块, 描述, 详情)


def test(name, condition, detail=""):
    """单个测试断言"""
    global passed, failed
    if condition:
        passed += 1
        print(f"  [PASS] {name}")
    else:
        failed += 1
        print(f"  [FAIL] {name}  {detail}")


def record_bug(module, desc, detail=""):
    """记录 bug"""
    bugs.append((module, desc, detail))
    print(f"  [BUG] [{module}] {desc}")
    if detail:
        print(f"        详情: {detail}")


# ==================== 1. Parser 测试 ====================

def test_parser():
    print("\n" + "=" * 60)
    print("1. Parser 智能解析测试")
    print("=" * 60)

    p = SmartParser()

    # 1.1 基础字段识别
    print("\n--- 1.1 基础字段识别 ---")
    r = p.parse("张三 360700199001011234 13812345678")
    d = p.to_variable_dict(r)
    test("姓名识别", d.get("姓名") == "张三", f"实际: {d.get('姓名')}")
    test("身份证识别", d.get("身份证号") == "360700199001011234", f"实际: {d.get('身份证号')}")
    test("手机号识别", d.get("手机号") == "13812345678", f"实际: {d.get('手机号')}")

    # 1.2 键值对识别
    print("\n--- 1.2 键值对识别 ---")
    r = p.parse("姓名：张三\n身份证号：360700199001011234\n手机号：13812345678")
    d = p.to_variable_dict(r)
    test("冒号键值对-姓名", d.get("姓名") == "张三", f"实际: {d.get('姓名')}")
    test("冒号键值对-身份证", d.get("身份证号") == "360700199001011234")

    r = p.parse("姓名:李四\n手机号:13900000000")
    d = p.to_variable_dict(r)
    test("英文冒号键值对", d.get("姓名") == "李四")

    r = p.parse("姓名=王五")
    d = p.to_variable_dict(r)
    test("等号键值对", d.get("姓名") == "王五", f"实际: {d.get('姓名')}")

    # 1.3 "是" 不再作为分隔符（避免"不是"等歧义）
    print("\n--- 1.3 '是'分隔符处理 ---")
    r = p.parse("姓名是赵六")
    d = p.to_variable_dict(r)
    # "是"已移除数，"姓名是赵六"应作为普通片段（5字不匹配姓名规则）
    test("'是'不作为分隔符", "姓名" not in d, f"实际: {d}")

    r = p.parse("姓名不是张三")
    d = p.to_variable_dict(r)
    test("'不是'不被误拆", "姓名不" not in d, f"实际: {d}")

    # 1.4 日期格式识别
    print("\n--- 1.4 日期格式识别 ---")
    for date_str, expected in [
        ("20260101", "2026-01-01"),
        ("2026-01-01", "2026-01-01"),
        ("2026/1/1", "2026-01-01"),
        ("2026年1月1日", "2026-01-01"),
    ]:
        r = p.parse(date_str)
        d = p.to_variable_dict(r)
        test(f"日期 {date_str} -> {expected}", d.get("日期") == expected,
             f"实际: {d.get('日期')}")

    # 1.5 日期范围识别
    print("\n--- 1.5 日期范围识别 ---")
    for text in ["20260101-20261030", "20260101~20261030"]:
        r = p.parse(text)
        d = p.to_variable_dict(r)
        test(f"范围 {text} 开始日期", d.get("开始日期") == "2026-01-01", f"实际: {d.get('开始日期')}")
        test(f"范围 {text} 结束日期", d.get("结束日期") == "2026-10-30", f"实际: {d.get('结束日期')}")

    # 检查日期范围是否还产生了多余的"日期"字段
    r = p.parse("20260101-20261030")
    fields = {f.field: f.value for f in r.fields}
    has_dup_date = "日期" in fields
    test("日期范围不产生多余日期字段", not has_dup_date,
         f"意外字段: {list(fields.keys())}")

    # 1.6 多人解析
    print("\n--- 1.6 多人解析 ---")
    text = "张三\n360700199001011234\n---\n李四\n360700199203035678"
    results = p.parse_multi(text)
    test("多人解析-人数", len(results) == 2, f"实际: {len(results)}")
    if len(results) == 2:
        d1 = p.to_variable_dict(results[0])
        d2 = p.to_variable_dict(results[1])
        test("第1人姓名", d1.get("姓名") == "张三")
        test("第2人姓名", d2.get("姓名") == "李四")

    # 测试 === 分隔符
    text2 = "王五\n360700199001011111\n===\n赵六\n360700199001012222"
    results2 = p.parse_multi(text2)
    test("===分隔符", len(results2) == 2, f"实际: {len(results2)}")

    # 1.7 空文本
    print("\n--- 1.7 空文本 ---")
    r = p.parse("")
    test("空文本不崩溃", r is not None and r.field_count == 0)

    r = p.parse(None)
    test("None不崩溃", r is not None)

    # 1.8 冲突检测
    print("\n--- 1.8 冲突检测 ---")
    r = p.parse("张三 李四")  # 两个姓名
    name_fields = [f for f in r.fields if f.field == "姓名"]
    if name_fields:
        test("姓名冲突标记", name_fields[0].conflict == True,
             f"conflict={name_fields[0].conflict}, alternatives={name_fields[0].alternatives}")

    # 1.9 单位识别
    print("\n--- 1.9 单位识别 ---")
    r = p.parse("某某科技有限公司")
    d = p.to_variable_dict(r)
    test("公司识别", d.get("单位") == "某某科技有限公司", f"实际: {d.get('单位')}")

    r = p.parse("某某银行")
    d = p.to_variable_dict(r)
    test("银行识别", d.get("单位") == "某某银行", f"实际: {d.get('单位')}")

    # 1.10 地址识别 + 拆分
    print("\n--- 1.10 地址识别与拆分 ---")
    r = p.parse("江西省赣州市章贡区红旗大道123号")
    d = p.to_variable_dict(r)
    test("地址识别", "地址" in d, f"实际字段: {list(d.keys())}")
    if "地址" in d:
        test("省识别", d.get("省") == "江西省", f"实际省: {d.get('省')}")
        test("市识别", d.get("市") == "赣州市", f"实际市: {d.get('市')}")
        test("区识别", d.get("区") == "章贡区", f"实际区: {d.get('区')}")

    # 1.11 未识别片段
    print("\n--- 1.11 未识别片段 ---")
    r = p.parse("张三 abc123 未知内容")
    test("未识别片段不为空", len(r.unrecognized) > 0,
         f"unrecognized: {r.unrecognized}")

    # 1.12 单行多字段（空格分隔）
    print("\n--- 1.12 单行多字段 ---")
    r = p.parse("张三 360700199001011234 13812345678 某某公司")
    d = p.to_variable_dict(r)
    test("单行姓名", d.get("姓名") == "张三")
    test("单行身份证", d.get("身份证号") == "360700199001011234")
    test("单行手机号", d.get("手机号") == "13812345678")
    test("单行单位", d.get("单位") == "某某公司", f"实际: {d.get('单位')}")


# ==================== 2. Scanner 测试 ====================

def test_scanner():
    print("\n" + "=" * 60)
    print("2. Scanner 模板扫描测试")
    print("=" * 60)

    businesses = scan_businesses(TEST_TEMPLATE_ROOT)
    test("扫描到业务", len(businesses) > 0, f"实际: {len(businesses)}")

    biz_names = [b.name for b in businesses]
    test("包含标准业务", "标准业务" in biz_names, f"实际: {biz_names}")

    # 找到标准业务
    std_biz = None
    for b in businesses:
        if b.name == "标准业务":
            std_biz = b
            break

    if std_biz:
        print(f"\n标准业务模板数: {len(std_biz.templates)}")
        print(f"合并变量: {sorted(std_biz.variables)}")

        test("变量包含姓名", "姓名" in std_biz.variables)
        test("变量包含身份证号", "身份证号" in std_biz.variables)

        # 检查每个模板的变量提取
        for tmpl in std_biz.templates:
            vars_count = len(tmpl.variables)
            test(f"模板 {tmpl.name} 有变量", vars_count > 0,
                 f"实际: {vars_count}")

    # 测试特殊字符业务
    spec_biz = None
    for b in businesses:
        if b.name == "特殊字符业务":
            spec_biz = b
            break
    if spec_biz:
        test("特殊字符业务有变量", len(spec_biz.variables) > 0)


# ==================== 3. Renderer 测试 ====================

def test_renderer():
    print("\n" + "=" * 60)
    print("3. Renderer 渲染测试")
    print("=" * 60)

    os.makedirs(TEST_OUTPUT_ROOT, exist_ok=True)

    # 3.1 标准 docx 渲染
    print("\n--- 3.1 docx 渲染 ---")
    tmpl = os.path.join(TEST_TEMPLATE_ROOT, "标准业务", "申请表.docx")
    out = os.path.join(TEST_OUTPUT_ROOT, "申请表_out.docx")
    variables = {
        "姓名": "张三",
        "身份证号": "360700199001011234",
        "手机号": "13812345678",
        "申请日期": "2026-10-16",
        "客户名": "张三",
        "单位": "某某科技有限公司",
        "地址": "北京市朝阳区建国路88号",
        "备注": "测试备注",
    }
    try:
        render_document(tmpl, out, variables)
        test("docx渲染成功", os.path.exists(out))

        # 验证内容
        import docx
        doc = docx.Document(out)
        full_text = "\n".join(p.text for p in doc.paragraphs)
        test("docx姓名替换", "张三" in full_text, f"文本: {full_text[:100]}")
        test("docx身份证替换", "360700199001011234" in full_text)
    except Exception as e:
        record_bug("renderer", f"docx渲染异常: {e}", traceback.format_exc())

    # 3.2 ${} 格式渲染
    print("\n--- 3.2 ${}格式渲染 ---")
    tmpl2 = os.path.join(TEST_TEMPLATE_ROOT, "标准业务", "美元符号模板.docx")
    out2 = os.path.join(TEST_OUTPUT_ROOT, "美元符号_out.docx")
    try:
        render_document(tmpl2, out2, variables)
        import docx
        doc = docx.Document(out2)
        full_text = "\n".join(p.text for p in doc.paragraphs)
        test("${}姓名替换", "张三" in full_text)
        test("${}无残留标记", "${" not in full_text and "}}" not in full_text,
             f"残留: {full_text}")
    except Exception as e:
        record_bug("renderer", "${}格式渲染异常: " + str(e))

    # 3.3 带空格的变量标记
    print("\n--- 3.3 带空格变量标记 ---")
    tmpl3 = os.path.join(TEST_TEMPLATE_ROOT, "标准业务", "带空格模板.docx")
    out3 = os.path.join(TEST_OUTPUT_ROOT, "带空格_out.docx")
    try:
        render_document(tmpl3, out3, variables)
        import docx
        doc = docx.Document(out3)
        full_text = "\n".join(p.text for p in doc.paragraphs)
        test("带空格替换", "张三" in full_text and "360700199001011234" in full_text)
    except Exception as e:
        record_bug("renderer", f"带空格渲染异常: {e}")

    # 3.4 跨 run 变量
    print("\n--- 3.4 跨run变量 ---")
    tmpl4 = os.path.join(TEST_TEMPLATE_ROOT, "标准业务", "跨run模板.docx")
    out4 = os.path.join(TEST_OUTPUT_ROOT, "跨run_out.docx")
    try:
        render_document(tmpl4, out4, variables)
        import docx
        doc = docx.Document(out4)
        full_text = "\n".join(p.text for p in doc.paragraphs)
        test("跨run替换", "张三" in full_text, f"实际: {full_text}")
        test("跨run无残留", "{{" not in full_text, f"残留: {full_text}")
    except Exception as e:
        record_bug("renderer", f"跨run渲染异常: {e}")

    # 3.5 页眉页脚变量
    print("\n--- 3.5 页眉页脚变量 ---")
    tmpl5 = os.path.join(TEST_TEMPLATE_ROOT, "标准业务", "页眉页脚模板.docx")
    out5 = os.path.join(TEST_OUTPUT_ROOT, "页眉页脚_out.docx")
    try:
        render_document(tmpl5, out5, variables)
        import docx
        doc = docx.Document(out5)
        header_text = "\n".join(p.text for p in doc.sections[0].header.paragraphs)
        footer_text = "\n".join(p.text for p in doc.sections[0].footer.paragraphs)
        test("页眉替换", "张三" in header_text, f"页眉: {header_text}")
        test("页脚替换", "360700199001011234" in footer_text, f"页脚: {footer_text}")
    except Exception as e:
        record_bug("renderer", f"页眉页脚渲染异常: {e}")

    # 3.6 xlsx 渲染
    print("\n--- 3.6 xlsx渲染 ---")
    tmpl6 = os.path.join(TEST_TEMPLATE_ROOT, "标准业务", "信息登记表.xlsx")
    out6 = os.path.join(TEST_OUTPUT_ROOT, "信息登记表_out.xlsx")
    try:
        render_document(tmpl6, out6, variables)
        import openpyxl
        wb = openpyxl.load_workbook(out6)
        ws = wb.active
        test("xlsx姓名替换", ws["B1"].value == "张三", f"实际: {ws['B1'].value}")
        test("xlsx身份证替换", ws["B2"].value == "360700199001011234")
        # ${} 格式
        test("xlsx美元符号替换", ws["B7"].value == "张三", f"实际: {ws['B7'].value}")
        wb.close()
    except Exception as e:
        record_bug("renderer", f"xlsx渲染异常: {e}")

    # 3.7 缺失变量（保留原标记）
    print("\n--- 3.7 缺失变量保留原标记 ---")
    tmpl7 = os.path.join(TEST_TEMPLATE_ROOT, "标准业务", "申请表.docx")
    out7 = os.path.join(TEST_OUTPUT_ROOT, "缺失变量_out.docx")
    partial_vars = {"姓名": "张三"}  # 只提供姓名
    try:
        render_document(tmpl7, out7, partial_vars)
        import docx
        doc = docx.Document(out7)
        full_text = "\n".join(p.text for p in doc.paragraphs)
        test("缺失变量保留标记", "{{身份证号}}" in full_text,
             f"文本: {full_text}")
    except Exception as e:
        record_bug("renderer", f"缺失变量渲染异常: {e}")

    # 3.8 特殊字符值
    print("\n--- 3.8 特殊字符值 ---")
    special_vars = {
        "姓名": "张三<>&\"'",
        "身份证号": "360700199001011234",
        "手机号": "13812345678",
        "申请日期": "2026-10-16",
        "客户名": "测试",
        "单位": "A&B公司",
        "地址": "路<号>",
        "备注": "备注'特殊\"字符",
    }
    out8 = os.path.join(TEST_OUTPUT_ROOT, "特殊字符_out.docx")
    try:
        render_document(tmpl, out8, special_vars)
        import docx
        doc = docx.Document(out8)
        full_text = "\n".join(p.text for p in doc.paragraphs)
        test("特殊字符不崩溃", True)
        test("特殊字符姓名", "张三<>&\"'" in full_text, f"实际: {full_text}")
    except Exception as e:
        record_bug("renderer", f"特殊字符渲染异常: {e}", traceback.format_exc())


# ==================== 4. Variable Manager 测试 ====================

def test_variable_manager():
    print("\n" + "=" * 60)
    print("4. Variable Manager 变量管理测试")
    print("=" * 60)

    vm = VariableManager("公务卡申请")
    vm.init_from_template_vars(["姓名", "身份证号", "手机号", "申请日期", "备注"])

    # 4.1 基本设置/获取
    print("\n--- 4.1 基本设置获取 ---")
    vm.set_value("姓名", "张三")
    test("设置值", vm.get_value("姓名") == "张三")

    # 4.2 默认表达式
    print("\n--- 4.2 默认表达式 ---")
    resolved = vm.get_all_resolved()
    test("申请日期默认today()", resolved.get("申请日期") != "",
         f"实际: {resolved.get('申请日期')}")
    import datetime
    today = datetime.date.today().strftime("%Y-%m-%d")
    test("申请日期=今天", resolved.get("申请日期") == today,
         f"期望: {today}, 实际: {resolved.get('申请日期')}")

    # 4.3 表达式求值
    print("\n--- 4.3 表达式求值 ---")
    vm.set_value("身份证号", "360700199001011234")
    # 手动测试 substr
    result = evaluate("substr(身份证号, 6, 8)", {"身份证号": "360700199001011234"})
    test("substr取出生日期", result == "19900101", f"实际: {result}")

    result = evaluate("concat(姓名, '-', 手机号)", {"姓名": "张三", "手机号": "138"})
    test("concat拼接", result == "张三-138", f"实际: {result}")

    # 4.4 链式依赖
    print("\n--- 4.4 链式依赖 ---")
    vm2 = VariableManager("test")
    vm2.init_from_template_vars(["A", "B", "C"])
    vm2.set_value("A", "10")
    vm2.variables["B"].value = "int(A) + 5"
    vm2.variables["C"].value = "int(B) * 2"
    resolved2 = vm2.get_all_resolved()
    test("链式A=10", resolved2.get("A") == "10")
    test("链式B=15", str(resolved2.get("B")) == "15", f"实际: {resolved2.get('B')}")
    test("链式C=30", str(resolved2.get("C")) == "30", f"实际: {resolved2.get('C')}")

    # 4.5 校验
    print("\n--- 4.5 校验 ---")
    vm3 = VariableManager("test")
    vm3.init_from_template_vars(["必填1", "必填2"])
    missing = vm3.validate()
    test("空值校验", len(missing) == 2, f"实际: {missing}")

    vm3.set_value("必填1", "值1")
    missing = vm3.validate()
    test("部分填充校验", missing == ["必填2"], f"实际: {missing}")

    # 4.6 CSV 导入导出
    print("\n--- 4.6 CSV导入导出 ---")
    vm4 = VariableManager("test")
    vm4.init_from_template_vars(["姓名", "身份证号"])
    csv_text = "姓名,张三\n身份证号,360700199001011234\n"
    count = vm4.load_from_csv(csv_text)
    test("CSV导入数量", count == 2, f"实际: {count}")
    test("CSV导入值", vm4.get_value("姓名") == "张三")

    exported = vm4.export_to_csv()
    test("CSV导出非空", len(exported) > 0)

    # 4.7 is_expression 边界
    print("\n--- 4.7 is_expression边界 ---")
    test("普通文本不是表达式", not is_expression("张三"))
    test("函数调用是表达式", is_expression("today()"))
    test("带括号文本", is_expression("测试(括号)"))  # 这个可能是 bug


# ==================== 5. Output 测试 ====================

def test_output():
    print("\n" + "=" * 60)
    print("5. Output 输出模块测试")
    print("=" * 60)

    from config.settings import get_settings
    settings = get_settings()
    original_output = settings.output_root_dir
    settings.output_root_dir = TEST_OUTPUT_ROOT

    om = OutputManager()

    # 5.1 子文件夹命名
    print("\n--- 5.1 子文件夹命名 ---")
    folder = om.build_subfolder_name("测试业务", {"客户名": "张三"})
    test("子文件夹命名", "测试业务" in folder and "张三" in folder,
         f"实际: {folder}")

    # 5.2 客户名为空
    folder2 = om.build_subfolder_name("测试业务", {})
    test("空客户名兜底", "未知客户" in folder2, f"实际: {folder2}")

    # 5.3 文件名非法字符清理
    print("\n--- 5.3 文件名清理 ---")
    folder3 = om.build_subfolder_name("测试", {"客户名": "张<三>:"})
    test("非法字符清理", "<" not in folder3 and ">" not in folder3 and ":" not in folder3,
         f"实际: {folder3}")

    # 5.4 重名策略 - append_seq
    print("\n--- 5.4 重名策略 ---")
    settings.rename_strategy = "append_seq"
    out_dir = om.get_output_dir("测试业务", {"客户名": "测试"})
    test_path = os.path.join(out_dir, "test.txt")
    with open(test_path, "w") as f:
        f.write("first")
    # 第二次应该追加序号
    path2 = om.resolve_output_path("test.txt", out_dir, {"身份证号": "360700199001011234"})
    test("追加序号不冲突", path2 != test_path, f"path2: {path2}")
    if path2:
        test("序号文件名含身份证后4位", "1234" in os.path.basename(path2),
             f"实际: {os.path.basename(path2)}")

    # 5.5 重名策略 - overwrite
    settings.rename_strategy = "overwrite"
    path3 = om.resolve_output_path("test.txt", out_dir, {})
    test("overwrite返回原路径", path3 == test_path, f"实际: {path3}")

    # 5.6 重名策略 - skip
    settings.rename_strategy = "skip"
    path4 = om.resolve_output_path("test.txt", out_dir, {})
    test("skip返回空", path4 == "", f"实际: {path4}")

    # 恢复
    settings.rename_strategy = "append_seq"
    settings.output_root_dir = original_output


# ==================== 6. 全流程端到端测试 ====================

def test_e2e():
    print("\n" + "=" * 60)
    print("6. 全流程端到端测试")
    print("=" * 60)

    from config.settings import get_settings
    settings = get_settings()
    original_template = settings.template_root_dir
    original_output = settings.output_root_dir
    settings.template_root_dir = TEST_TEMPLATE_ROOT
    settings.output_root_dir = TEST_OUTPUT_ROOT

    # 6.1 单人全流程
    print("\n--- 6.1 单人全流程 ---")
    businesses = scan_businesses(TEST_TEMPLATE_ROOT)
    test("扫描业务", len(businesses) > 0)

    biz = None
    for b in businesses:
        if b.name == "标准业务":
            biz = b
            break

    if biz:
        from core.variable_manager import VariableManager
        vm = VariableManager(biz.name)
        vm.init_from_template_vars(sorted(biz.variables))

        # 模拟智能粘贴（用键值对确保字段名匹配模板）
        p = SmartParser()
        r = p.parse("姓名：张三\n身份证号：360700199001011234\n手机号：13812345678\n单位：某某科技有限公司\n申请日期：20261016")
        d = p.to_variable_dict(r)
        for k, v in d.items():
            if k in vm.variables:
                vm.set_value(k, v)
        # 补充其他必填变量
        vm.set_value("客户名", "张三")
        vm.set_value("出生日期", "1990-01-01")
        vm.set_value("地址", "江西省赣州市章贡区红旗大道123号")
        vm.set_value("备注", "测试")
        vm.set_value("开始日期", "2026-01-01")
        vm.set_value("结束日期", "2026-10-30")

        # 校验
        missing = vm.validate()
        test("全流程校验通过", len(missing) == 0, f"缺失: {missing}")

        resolved = vm.get_all_resolved()
        test("解析值非空", len(resolved) > 0)

        # 生成所有文档
        om = OutputManager()
        out_dir = om.get_output_dir(biz.name, resolved)
        success = 0
        fail = 0
        for tmpl in biz.templates:
            try:
                out_path = om.resolve_output_path(tmpl.rel_path, out_dir, resolved)
                if out_path:
                    render_document(tmpl.path, out_path, resolved)
                    success += 1
            except Exception as e:
                fail += 1
                record_bug("e2e", f"模板 {tmpl.name} 生成失败: {e}")

        test(f"全流程生成成功 {success}/{len(biz.templates)}",
             success > 0, f"成功: {success}, 失败: {fail}")

    # 6.2 多人批量全流程
    print("\n--- 6.2 多人批量全流程 ---")
    multi_text = "张三 360700199001011234 13812345678\n---\n李四 360700199203035678 13987654321"
    p = SmartParser()
    results = p.parse_multi(multi_text)
    test("多人解析", len(results) == 2)

    if biz and len(results) == 2:
        om = OutputManager()
        total_success = 0
        for person_r in results:
            d = p.to_variable_dict(person_r)
            vm_p = VariableManager(biz.name)
            vm_p.init_from_template_vars(sorted(biz.variables))
            for k, v in d.items():
                if k in vm_p.variables:
                    vm_p.set_value(k, v)
            resolved = vm_p.get_all_resolved()
            out_dir = om.get_output_dir(biz.name, resolved)
            for tmpl in biz.templates:
                try:
                    out_path = om.resolve_output_path(tmpl.rel_path, out_dir, resolved)
                    if out_path:
                        render_document(tmpl.path, out_path, resolved)
                        total_success += 1
                except Exception as e:
                    record_bug("e2e_multi", f"批量生成失败: {e}")

        test(f"多人生成 {total_success} 个文档", total_success > 0)

    # 恢复设置
    settings.template_root_dir = original_template
    settings.output_root_dir = original_output


# ==================== 7. 边界/压力测试 ====================

def test_edge_cases():
    print("\n" + "=" * 60)
    print("7. 边界/压力测试")
    print("=" * 60)

    p = SmartParser()

    # 7.1 超长文本
    print("\n--- 7.1 超长文本 ---")
    long_text = "张三 " * 1000 + "360700199001011234 " * 100
    try:
        r = p.parse(long_text)
        test("超长文本不崩溃", r is not None)
    except Exception as e:
        record_bug("parser", f"超长文本崩溃: {e}", traceback.format_exc())

    # 7.2 纯空白
    print("\n--- 7.2 纯空白 ---")
    r = p.parse("   \n  \t  \n")
    test("纯空白不崩溃", r is not None and r.field_count == 0)

    # 7.3 特殊字符输入
    print("\n--- 7.3 特殊字符输入 ---")
    r = p.parse("张三\n360700199001011234\n<script>alert(1)</script>")
    test("特殊字符输入不崩溃", r is not None)
    test("脚本未被识别为字段", "script" not in p.to_variable_dict(r).get("姓名", ""))

    # 7.4 身份证末尾X
    print("\n--- 7.4 身份证末尾X ---")
    r = p.parse("36070019900101123X")
    d = p.to_variable_dict(r)
    test("身份证末尾X", d.get("身份证号") == "36070019900101123X",
         f"实际: {d.get('身份证号')}")

    # 7.5 日期范围中文分隔
    print("\n--- 7.5 日期范围中文分隔 ---")
    r = p.parse("2026年1月1日至2026年10月30日")
    d = p.to_variable_dict(r)
    test("中文至分隔-开始", d.get("开始日期") == "2026-01-01", f"实际: {d.get('开始日期')}")
    test("中文至分隔-结束", d.get("结束日期") == "2026-10-30", f"实际: {d.get('结束日期')}")

    # 7.6 多人无分隔符（单人模式）
    print("\n--- 7.6 单人数据不误判为多人 ---")
    text = "张三\n360700199001011234\n13812345678"
    results = p.parse_multi(text)
    test("单人不拆分", len(results) == 1, f"实际: {len(results)}")

    # 7.7 连续分隔符
    print("\n--- 7.7 连续分隔符 ---")
    text = "张三\n---\n---\n李四"
    results = p.parse_multi(text)
    test("连续分隔符", len(results) == 2, f"实际: {len(results)}")

    # 7.8 数字 0 作为值
    print("\n--- 7.8 数字0作为值 ---")
    vm = VariableManager("test")
    vm.init_from_template_vars(["数量"])
    vm.set_value("数量", 0)
    missing = vm.validate()
    test("数字0不被当作空", len(missing) == 0, f"实际缺失: {missing}")

    # 7.9 变量名包含特殊字符
    print("\n--- 7.9 变量名特殊字符 ---")
    from core.scanner import extract_variables_from_text
    vars1 = extract_variables_from_text("{{变量-名}} ${变量_名}")
    test("变量名含-", "变量-名" in vars1, f"实际: {vars1}")
    test("变量名含_", "变量_名" in vars1, f"实际: {vars1}")

    # 7.10 空变量值渲染
    print("\n--- 7.10 空变量值渲染 ---")
    tmpl = os.path.join(TEST_TEMPLATE_ROOT, "标准业务", "申请表.docx")
    out = os.path.join(TEST_OUTPUT_ROOT, "空值渲染.docx")
    vars_empty = {"姓名": "", "身份证号": "", "手机号": "", "申请日期": "",
                  "客户名": "", "单位": "", "地址": "", "备注": ""}
    try:
        render_document(tmpl, out, vars_empty)
        test("空值渲染不崩溃", os.path.exists(out))
    except Exception as e:
        record_bug("renderer", f"空值渲染崩溃: {e}")


def main():
    print("=" * 60)
    print("全面测试开始")
    print("=" * 60)

    test_parser()
    test_scanner()
    test_renderer()
    test_variable_manager()
    test_output()
    test_e2e()
    test_edge_cases()

    print("\n" + "=" * 60)
    print(f"测试结果汇总：通过 {passed}，失败 {failed}，发现 Bug {len(bugs)}")
    print("=" * 60)

    if bugs:
        print("\n--- Bug 列表 ---")
        for i, (module, desc, detail) in enumerate(bugs, 1):
            print(f"\n[{i}] [{module}] {desc}")
            if detail:
                print(f"    {detail[:200]}")

    return len(bugs) == 0 and failed == 0


if __name__ == "__main__":
    success = main()
    sys.exit(0 if success else 1)
