"""
创建测试用模板文件，验证全流程。
生成：
  模板根目录/
    公务卡申请/
      申请表.docx
      信息登记表.xlsx
"""

import os
import sys

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from config.settings import get_settings

TEST_TEMPLATE_ROOT = os.path.join(os.path.dirname(os.path.abspath(__file__)), "test_templates")


def create_docx_template():
    """生成测试用 .docx 模板，包含段落变量和表格变量"""
    import docx
    from docx.shared import Pt

    path = os.path.join(TEST_TEMPLATE_ROOT, "公务卡申请", "申请表.docx")
    os.makedirs(os.path.dirname(path), exist_ok=True)

    doc = docx.Document()
    doc.add_heading("公务卡申请表", level=1)
    doc.add_paragraph(f"姓名：{{{{姓名}}}}")
    doc.add_paragraph(f"身份证号：{{{{身份证号}}}}")
    doc.add_paragraph(f"手机号：{{{{手机号}}}}")
    doc.add_paragraph(f"申请日期：{{{{申请日期}}}}")

    # 表格中也放变量
    table = doc.add_table(rows=3, cols=2)
    table.style = "Table Grid"
    table.cell(0, 0).text = "客户名"
    table.cell(0, 1).text = "{{客户名}}"
    table.cell(1, 0).text = "单位"
    table.cell(1, 1).text = "{{单位}}"
    table.cell(2, 0).text = "出生日期"
    table.cell(2, 1).text = "{{出生日期}}"

    doc.save(path)
    print(f"创建：{path}")


def create_xlsx_template():
    """生成测试用 .xlsx 模板，单元格中包含变量标记"""
    import openpyxl

    path = os.path.join(TEST_TEMPLATE_ROOT, "公务卡申请", "信息登记表.xlsx")
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "信息登记"
    ws["A1"] = "姓名"
    ws["B1"] = "{{姓名}}"
    ws["A2"] = "身份证号"
    ws["B2"] = "{{身份证号}}"
    ws["A3"] = "手机号"
    ws["B3"] = "{{手机号}}"
    ws["A4"] = "单位"
    ws["B4"] = "{{单位}}"
    ws["A5"] = "申请日期"
    ws["B5"] = "{{申请日期}}"
    wb.save(path)
    print(f"创建：{path}")


def main():
    """生成所有测试模板并将 tests/test_templates 设为模板根目录"""
    create_docx_template()
    create_xlsx_template()
    print(f"\n测试模板根目录：{TEST_TEMPLATE_ROOT}")
    # 设置为模板根目录
    settings = get_settings()
    settings.template_root_dir = TEST_TEMPLATE_ROOT
    print(f"已设置模板根目录：{settings.template_root_dir}")


if __name__ == "__main__":
    main()
