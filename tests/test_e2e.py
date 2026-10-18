"""
端到端测试：扫描 -> 变量填充 -> 渲染 -> 输出
直接调用各模块，验证阶段一全流程。
"""

import os
import sys

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from config.settings import get_settings
from core.scanner import scan_businesses
from core.variable_manager import VariableManager
from core.renderer import render_document, RenderError
from output.file_output import OutputManager, OutputError


def main():
    """端到端测试：扫描 -> 填充变量 -> 渲染 -> 输出，验证全流程"""
    settings = get_settings()
    template_root = settings.template_root_dir
    output_root = os.path.join(os.path.dirname(os.path.abspath(__file__)), "test_output")
    settings.output_root_dir = output_root

    print(f"模板根目录：{template_root}")
    print(f"输出根目录：{output_root}")

    # 1. 扫描
    businesses = scan_businesses(template_root)
    print(f"\n扫描到 {len(businesses)} 项业务")
    for b in businesses:
        print(f"  - {b.name}: 模板 {len(b.templates)} 个, 变量 {len(b.variables)} 个")
        for v in sorted(b.variables):
            print(f"      变量: {v}")

    business = businesses[0]

    # 2. 初始化变量
    vm = VariableManager(business_name=business.name)
    vm.init_from_template_vars(sorted(business.variables))

    # 3. 填充变量值
    test_values = {
        "姓名": "张三",
        "身份证号": "360700199001011234",
        "手机号": "13812345678",
        "客户名": "张三",
        "单位": "某某科技有限公司",
        "出生日期": "substr(身份证号, 6, 8)",  # 表达式：从身份证取出生日期
        "申请日期": "",  # 走默认值 today()
    }
    for name, value in test_values.items():
        vm.set_value(name, value)

    # 校验
    missing = vm.validate()
    if missing:
        print(f"\n[校验失败] 必填变量缺失：{missing}")
        return
    print("\n[校验通过] 必填变量已全部填写")

    # 解析值
    resolved = vm.get_all_resolved()
    print("\n解析后的变量值：")
    for k, v in resolved.items():
        print(f"  {k} = {v}")

    # 4. 输出
    output_mgr = OutputManager()
    try:
        output_dir = output_mgr.get_output_dir(business.name, resolved)
    except OutputError as e:
        print(f"[输出错误] {e}")
        return

    print(f"\n输出目录：{output_dir}")

    # 5. 渲染每个模板
    success = 0
    fail = 0
    for tmpl in business.templates:
        if tmpl.error:
            print(f"  [跳过] {tmpl.name}: {tmpl.error}")
            fail += 1
            continue
        output_path = output_mgr.resolve_output_path(tmpl.rel_path, output_dir, resolved)
        if not output_path:
            print(f"  [跳过] {tmpl.name}: 文件已存在")
            continue
        try:
            render_document(tmpl.path, output_path, resolved)
            print(f"  [成功] {tmpl.name} -> {os.path.basename(output_path)}")
            success += 1
        except RenderError as e:
            print(f"  [失败] {tmpl.name}: {e}")
            fail += 1

    print(f"\n完成：成功 {success}，失败 {fail}")

    # 验证输出文件内容
    print("\n--- 验证输出文件 ---")
    for f in os.listdir(output_dir):
        fpath = os.path.join(output_dir, f)
        if f.endswith(".docx"):
            import docx
            doc = docx.Document(fpath)
            print(f"\n[{f}] 段落内容：")
            for p in doc.paragraphs:
                if p.text.strip():
                    print(f"  {p.text}")
            for table in doc.tables:
                for row in table.rows:
                    print(f"  {' | '.join(c.text for c in row.cells)}")
        elif f.endswith(".xlsx"):
            import openpyxl
            wb = openpyxl.load_workbook(fpath)
            ws = wb.active
            print(f"\n[{f}] 单元格内容：")
            for row in ws.iter_rows(values_only=True):
                print(f"  {row}")
            wb.close()


if __name__ == "__main__":
    main()
