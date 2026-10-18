"""Fast smoke tests for the core pipeline.

The exhaustive script-based suite lives in ``comprehensive_test.py``.
Run this file with ``pytest``; run the full suite with
``python tests/comprehensive_test.py``.
"""

from pathlib import Path

import docx

from core.expression import evaluate
from core.parser import SmartParser
from core.renderer import render_document
from core.scanner import scan_businesses

ROOT = Path(__file__).resolve().parents[1]
TEMPLATES = ROOT / "examples" / "templates"


def test_parser_recognizes_common_fields():
    result = SmartParser().parse(
        "姓名：张三\n"
        "身份证号：360700199001011234\n"
        "手机号：13812345678\n"
        "地址：江西省赣州市章贡区红旗大道123号"
    )
    values = {f.field: f.value for f in result.fields}

    assert values["姓名"] == "张三"
    assert values["身份证号"] == "360700199001011234"
    assert values["手机号"] == "13812345678"
    assert values["省"] == "江西省"
    assert values["市"] == "赣州市"
    assert values["区"] == "章贡区"

    order = [f.field for f in result.fields]
    assert order.index("省") < order.index("市") < order.index("区")


def test_expression_engine():
    assert evaluate("substr(身份证号, 6, 8)", {"身份证号": "360700199001011234"}) == "19900101"
    assert evaluate("concat(姓, 名)", {"姓": "张", "名": "三"}) == "张三"
    assert evaluate("today()")


def test_scan_examples():
    businesses = scan_businesses(str(TEMPLATES))
    names = {b.name for b in businesses}

    assert "公务卡申请" in names
    assert "信用卡激活" in names

    application = next(b for b in businesses if b.name == "公务卡申请")
    assert "姓名" in application.variables
    assert len(application.templates) == 2


def test_render_docx(tmp_path):
    template = TEMPLATES / "公务卡申请" / "申请表.docx"
    output = tmp_path / "申请表_张三.docx"

    render_document(
        str(template),
        str(output),
        {
            "姓名": "张三",
            "身份证号": "360700199001011234",
            "手机号": "13812345678",
            "单位": "某某科技有限公司",
            "地址": "江西省赣州市章贡区红旗大道123号",
            "出生日期": "1990-01-01",
            "申请日期": "2026-10-18",
        },
    )

    document = docx.Document(str(output))
    text = "\n".join(p.text for p in document.paragraphs)
    for table in document.tables:
        for row in table.rows:
            text += "\n" + " | ".join(cell.text for cell in row.cells)

    assert "张三" in text
    assert "360700199001011234" in text
    assert "{{" not in text
    assert template.exists()
