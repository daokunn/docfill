"""P2 回归测试：目录防护、页眉表格、图标资源与新手教程。"""

from pathlib import Path

import docx
from docx.shared import Inches

from app_paths import is_subpath
from core.renderer import render_document
from core.scanner import extract_template_variables
from ui.style import TUTORIAL_HTML

ROOT = Path(__file__).resolve().parents[1]


def test_output_must_not_be_inside_template(tmp_path):
    template_root = tmp_path / "模板"
    template_root.mkdir()
    inside = template_root / "生成结果"
    sibling = tmp_path / "生成结果"

    assert is_subpath(str(inside), str(template_root)) is True
    assert is_subpath(str(template_root), str(template_root)) is True
    assert is_subpath(str(sibling), str(template_root)) is False


def test_header_table_variables_are_scanned(tmp_path):
    template = tmp_path / "header_table.docx"
    document = docx.Document()
    document.add_paragraph("正文：{{姓名}}")
    header_table = document.sections[0].header.add_table(rows=1, cols=2, width=Inches(6))
    header_table.cell(0, 0).text = "页眉变量"
    header_table.cell(0, 1).text = "{{页眉编号}}"
    document.save(template)

    variables = extract_template_variables(str(template))
    assert "姓名" in variables
    assert "页眉编号" in variables

    output = tmp_path / "header_table_out.docx"
    render_document(str(template), str(output), {"姓名": "张三", "页眉编号": "A-001"})
    rendered = docx.Document(output)
    header_text = "\n".join(
        cell.text for table in rendered.sections[0].header.tables
        for row in table.rows for cell in row.cells
    )
    assert "A-001" in header_text


def test_tutorial_is_beginner_friendly():
    assert "第一次用" in TUTORIAL_HTML
    assert "看这三步就够了" in TUTORIAL_HTML
    assert "黄色格子" in TUTORIAL_HTML
    assert "5329532" not in TUTORIAL_HTML


def test_brand_assets_exist():
    assert (ROOT / "assets" / "docfill.ico").is_file()
    assert (ROOT / "assets" / "docfill.png").is_file()
    assert (ROOT / "assets" / "splash.png").is_file()
