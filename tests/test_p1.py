"""P1 回归测试：可选变量、排序、缓存与 CSV。"""

from pathlib import Path

import docx

from core.csv_io import detect_header, parse_csv_rows, read_csv_file, rows_to_mapping
from core.preflight import collect_empty_variables
from core.renderer import _replace_in_text
from core.scanner import (
    BusinessInfo,
    TemplateInfo,
    extract_template_variable_specs,
    extract_template_variables,
    ordered_variable_names,
    scan_businesses,
)
from core.user_state import clear_last_data, load_last_data, save_last_data
from core.variable_manager import VariableManager

ROOT = Path(__file__).resolve().parents[1]
TEMPLATES = ROOT / "examples" / "templates"


def test_optional_variable_syntax(tmp_path):
    template = tmp_path / "optional.docx"
    document = docx.Document()
    document.add_paragraph("姓名：{{姓名}}")
    document.add_paragraph("备注：{{备注?}}")
    document.save(template)

    specs = extract_template_variable_specs(str(template))
    assert specs == [("姓名", False), ("备注", True)]
    assert extract_template_variables(str(template)) == {"姓名", "备注"}

    # 渲染时 ? 只是声明标记，不参与变量名
    assert _replace_in_text("{{备注?}}", {"备注": "已填写"}) == "已填写"


def test_variable_manager_optional():
    vm = VariableManager()
    vm.init_from_template_vars(["姓名", "备注"], {"备注"})

    assert vm.variables["姓名"].required is True
    assert vm.variables["备注"].required is False
    assert vm.validate() == ["姓名"]


def test_preflight_ignores_optional():
    business = BusinessInfo("测试业务", ".")
    template = TemplateInfo("测试.docx", "测试.docx", "测试.docx")
    template.variables = {"姓名", "备注"}
    business.templates = [template]
    business.variables = {"姓名", "备注"}
    business.optional_variables = {"备注"}

    assert collect_empty_variables(business, {"姓名": "张三"}) == []


def test_variable_order_setting():
    business = next(b for b in scan_businesses(str(TEMPLATES)) if b.name == "公务卡申请")

    template_order = ordered_variable_names(business, "template")
    alphabetical = ordered_variable_names(business, "alphabetical")

    assert template_order[:3] == ["姓名", "身份证号", "手机号"]
    assert template_order != alphabetical
    assert alphabetical == sorted(alphabetical)


def test_last_data_persistence(tmp_path):
    state_path = str(tmp_path / "last_data.json")

    save_last_data("公务卡申请", {"姓名": "张三", "备注": "测试"}, state_path)
    assert load_last_data("公务卡申请", state_path)["姓名"] == "张三"

    clear_last_data("公务卡申请", state_path)
    assert load_last_data("公务卡申请", state_path) == {}


def test_csv_encoding_and_header(tmp_path):
    path = tmp_path / "gbk.csv"
    path.write_bytes("变量名,值\n姓名,张三\n手机号,13812345678\n".encode("gbk"))

    text, encoding = read_csv_file(str(path))
    rows = parse_csv_rows(text)

    assert encoding in ("gb18030", "gbk")
    assert "姓名" in text
    assert detect_header(rows) is True
    assert rows_to_mapping(rows, has_header=True)["姓名"] == "张三"
    assert rows_to_mapping(rows, has_header=False)["变量名"] == "值"
