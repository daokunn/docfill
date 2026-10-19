"""溢出自动缩小字号的回归测试。"""

import zipfile

from docx import Document
from docx.shared import Pt
from lxml import etree

from core.renderer import render_document

W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"


def _make_narrow_cell_doc(path):
    document = Document()
    table = document.add_table(rows=1, cols=2)
    table.autofit = False
    # 固定列宽（dxa），让宽度可预测：约 4cm ≈ 2268 dxa
    row = table.rows[0]
    row.cells[0].text = "地址"
    para = row.cells[1].paragraphs[0]
    para.text = ""
    run = para.add_run("{{地址}}")
    run.font.size = Pt(14)
    from docx.shared import Cm
    row.cells[0].width = Cm(3)
    row.cells[1].width = Cm(4)
    document.save(path)


def _cell_font_half_pt(path):
    """读取表格第二格第一个 run 的字号（半磅）。"""
    data = zipfile.ZipFile(path).read("word/document.xml")
    root = etree.fromstring(data)
    cells = list(root.iter(f"{{{W}}}tc"))
    if len(cells) < 2:
        return None
    second_cell = cells[1]
    for run in second_cell.iter(f"{{{W}}}r"):
        r_pr = run.find(f"{{{W}}}rPr")
        if r_pr is not None:
            sz = r_pr.find(f"{{{W}}}sz")
            if sz is not None:
                return int(sz.get(f"{{{W}}}val"))
    return None


def test_long_value_triggers_shrink(tmp_path):
    template = tmp_path / "fit.docx"
    output = tmp_path / "fit_out.docx"
    _make_narrow_cell_doc(template)

    before_sz = _cell_font_half_pt(template)
    assert before_sz == 28  # 14pt

    log = []
    render_document(
        str(template), str(output),
        {"地址": "浙江省某某市某某区某某某某某某某某某某某某某某某某某某路123号"},
        shrink_overflow_font=True, layout_log=log,
    )
    after_sz = _cell_font_half_pt(output)
    assert after_sz is not None
    assert after_sz < before_sz, f"预期缩小，实际 {before_sz}→{after_sz}"
    assert after_sz >= int(before_sz * 0.7), "不应低于缩放下限"
    assert any("磅" in msg for msg in log), f"应有版式日志，实际 {log}"


def test_short_value_does_not_touch_format(tmp_path):
    template = tmp_path / "fit2.docx"
    output = tmp_path / "fit2_out.docx"
    _make_narrow_cell_doc(template)

    before = _cell_font_half_pt(template)
    log = []
    render_document(
        str(template), str(output), {"地址": "甲路1号"},
        shrink_overflow_font=True, layout_log=log,
    )
    assert _cell_font_half_pt(output) == before
    assert log == [], f"短内容不应触发缩小，实际 {log}"


def test_shrink_disabled_keeps_everything(tmp_path):
    template = tmp_path / "fit3.docx"
    output = tmp_path / "fit3_out.docx"
    _make_narrow_cell_doc(template)

    before = _cell_font_half_pt(template)
    log = []
    render_document(
        str(template), str(output),
        {"地址": "浙江省某某市某某区某某某某某某某某某某某某某某某某某某路123号"},
        shrink_overflow_font=False, layout_log=log,
    )
    assert _cell_font_half_pt(output) == before, "关闭开关后绝不能动字号"
    assert log == []
